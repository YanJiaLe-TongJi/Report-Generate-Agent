import io
import json
import queue
from pathlib import Path
from types import SimpleNamespace
import pytest
from app import create_app
from harness import Harness, execute, render
from paths import DATA_DIR
from storage import write_json
import providers

PROFILE={'provider':'custom','base_url':'http://localhost:9999/v1','model':'test'}
@pytest.fixture
def client():
    app=create_app('test-token');app.testing=True
    client=app.test_client();client.get('/?launch=test-token')
    return client

def post(client,path,data):return client.post(path,json=data,headers={'X-Desktop-Request':'1'})

def test_loopback_auth_and_csrf():
    app=create_app('secret');c=app.test_client()
    assert c.get('/api/state').status_code==403
    r=c.get('/?launch=secret');assert r.status_code==200
    assert 'HttpOnly' in r.headers['Set-Cookie']
    assert c.get('/api/state').status_code==200
    assert c.post('/api/stop').status_code==403
    assert c.get('/api/state',headers={'Host':'attacker.invalid'}).status_code==403
    assert c.post('/api/stop',headers={'X-Desktop-Request':'1','Origin':'https://attacker.invalid'}).status_code==403

def test_no_account_routes(client):
    for path in ['/login','/register','/admin','/api/user/profile','/api/experiment-assistant']:
        assert client.get(path).status_code==404

def test_credentials_not_serialized(client,monkeypatch):
    vault={}
    monkeypatch.setattr(providers.keyring,'set_password',lambda s,k,v:vault.update({k:v}))
    monkeypatch.setattr(providers.keyring,'get_password',lambda s,k:vault.get(k))
    r=post(client,'/api/settings',{'text':dict(PROFILE,api_key='sk-test-private')})
    assert r.status_code==200
    assert 'sk-test-private' not in (DATA_DIR/'settings.json').read_text()
    assert 'sk-test-private' not in client.get('/api/state').text
    assert providers.secret(PROFILE)=='sk-test-private'

def test_library_unicode_and_reuse(client):
    r=client.post('/api/library',data={'category':'data','files':(io.BytesIO('电压,电流\n1,2'.encode()),'实验 数据.csv')},headers={'X-Desktop-Request':'1'})
    assert r.status_code==200
    item=r.json['items'][0];assert item['name']=='实验 数据.csv'
    assert Path(item['path']).is_file()
    client.delete('/api/library/'+item['id'],headers={'X-Desktop-Request':'1'})
    assert Path(item['path']).is_file()  # old revisions retain their source files

def test_light_edition_rejects_pdf():
    with pytest.raises(ValueError,match='PDF'):Harness().start({'format_type':'latex'})

def test_restart_interrupts_pending():
    write_json(DATA_DIR/'tasks/restart.json',{'id':'restart','status':'processing','created':0})
    assert Harness().tasks['restart']['status']=='interrupted'

def test_parse_and_formula(tmp_path):
    from ai_generator import parse_excel_data,build_raw_data_confirmation_workbook
    from word_backend import create_omml_formula
    f=tmp_path/'中文 数据.csv';f.write_text('电压,电流\n1,2\n2,4',encoding='utf-8')
    parsed,diagnostics=parse_excel_data([f]);assert parsed[f.name]['CSV'][1]==['1','2']
    assert not diagnostics['failed_files']
    assert create_omml_formula(r'\frac{x^2}{\sqrt{y}}') is not None

def sample_task(tmp_path):
    from constants import KNOWN_SECTIONS
    return {'id':'sample','status':'processing','steps':[],'created':0,'config':{'task_dir':str(tmp_path),'cover_info':{'experiment_name':'伏安法测量电阻','student_name':'测试学生'},'format_type':'word','material_paths':[],'raw_data_paths':[],'data_paths':[],'example_paths':[],'text_profile':PROFILE,'text_model':'test'},'section_contents':{s:'这是实验内容。电阻满足 $R=U/I$。\n电压 | 电流\n1 | 2\n2 | 4' for s in KNOWN_SECTIONS}}

def test_word_render_without_office(tmp_path):
    from docx import Document
    task=sample_task(tmp_path);render(task)
    doc=Document(DATA_DIR/'outputs'/task['artifacts'][0])
    assert any('伏安法' in p.text for p in doc.paragraphs)
    assert doc.tables
    assert 'sk-' not in json.dumps(task)

def test_revision_preserves_original(tmp_path,monkeypatch):
    import harness
    task=sample_task(tmp_path);original=dict(task['section_contents'])
    task.update(instruction='补充结论',targets=['实验结论'])
    response=SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps({'实验结论':'新的实验结论'})))])
    fake=SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **kw:response)))
    monkeypatch.setattr(harness,'make_client',lambda p:fake)
    q=queue.Queue();execute(task,q)
    result=q.get();assert result['status']=='done'
    assert result['section_contents']['实验结论']=='新的实验结论'
    assert result['section_contents']['实验原理']==original['实验原理']

def test_provider_errors_redact_secrets(monkeypatch):
    def fail(**kw):raise RuntimeError('sk-secret-value https://private')
    monkeypatch.setattr(providers,'OpenAI',lambda **kw:SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=fail))))
    c=providers.Client(PROFILE,'sk-secret-value')
    with pytest.raises(RuntimeError) as exc:c.create(model='test',messages=[])
    assert 'sk-secret' not in str(exc.value)

@pytest.mark.parametrize('provider',['deepseek','kimi','qwen','minimax','glm','custom'])
def test_provider_adapter(provider,monkeypatch):
    calls=[]
    def complete(**kw):
        calls.append(kw);return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='OK'),finish_reason='stop')])
    monkeypatch.setattr(providers,'OpenAI',lambda **kw:SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=complete))))
    p=providers.PRESETS[provider];profile={'provider':provider,'base_url':p['base_url'] or PROFILE['base_url'],'model':p['text_model'] or 'test'}
    providers.Client(profile,'test').create(model=profile['model'],messages=[],max_tokens=16384)
    assert calls[0]['model']==profile['model']
    if provider=='custom':assert 'extra_body' not in calls[0]

def test_vision_requires_actual_reading(monkeypatch):
    monkeypatch.setattr(providers,'Client',lambda *a,**kw:SimpleNamespace(create=lambda **k:SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='OK'))])))
    with pytest.raises(ValueError,match='图片识别'):providers.test_connection(PROFILE,vision=True)

def test_stop_terminates_task_process():
    import multiprocessing as mp
    import time
    engine=Harness();p=mp.get_context('spawn').Process(target=time.sleep,args=(60,));p.start()
    engine.process=p;engine.active='stop-test';engine.tasks['stop-test']={'id':'stop-test','status':'processing','created':0}
    engine.stop()
    assert not p.is_alive()
    assert engine.tasks['stop-test']['status']=='interrupted'
    assert engine.active is None

def test_pdf_failure_keeps_source_and_log(tmp_path,monkeypatch):
    import harness,latex_backend
    monkeypatch.setattr(harness,'EDITION','full')
    monkeypatch.setattr(latex_backend,'compile_latex',lambda *a:(False,'compiler diagnostic'))
    task=sample_task(tmp_path);task['id']='failed-pdf';task['config']['format_type']='latex'
    with pytest.raises(ValueError,match='编译失败'):render(task)
    assert any(n.endswith('.zip') for n in task['artifacts'])
    assert any(n.endswith('.log') for n in task['artifacts'])
    assert any(n.endswith('.md') for n in task['artifacts'])
